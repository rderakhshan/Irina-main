"""Sandbox-side implementation for Irina's extended tools.

The trusted host only builds tool schemas and forwards one JSON request at a
time. Every operation in this module runs inside the approved OpenShell
sandbox: files are resolved below the workspace, network calls are subject to
the sandbox policy, and external programs run with the sandbox user.

Design constraints:
* No operation accepts a host path or opens a host database.
* Third-party libraries are imported lazily so a missing dependency produces a
  clear, model-readable error instead of breaking the whole tool surface.
* Results are deliberately compact: bounded rows, bounded text, and short
  diagnostics.
"""

from __future__ import annotations

import datetime as _dt
import html
import ipaddress
import json
import math
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

MAX_TEXT_CHARS = 20_000
MAX_HTTP_BYTES = 2_000_000
MAX_DOWNLOAD_BYTES = 50 * 1024 * 1024
MAX_ROWS = 500
USER_AGENT = "Irina/0.1 (+https://github.com/rderakhshan/Irina-main)"


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------
class ToolError(RuntimeError):
    """A user-visible operation failure."""


def _resolve(workdir: str | Path, path: str) -> Path:
    """Resolve a workspace-relative path and reject escapes."""
    root = Path(workdir).resolve()
    candidate = Path(str(path)).expanduser()
    if not candidate.is_absolute():
        candidate = root / candidate
    resolved = candidate.resolve()
    try:
        resolved.relative_to(root)
    except ValueError:
        raise PermissionError(
            f"{path!r} escapes the working directory"
        ) from None
    return resolved


def _clip(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    if len(text) <= limit:
        return text
    half = max(1, limit // 2)
    return text[:half] + "\n... [truncated] ...\n" + text[-half:]


def _jsonable(value: Any) -> Any:
    """Convert pandas/numpy/sqlite values into JSON-safe objects."""
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return None if math.isnan(value) or math.isinf(value) else value
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", "replace")
    if isinstance(value, dict):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(item) for item in value]
    try:
        import numpy as np  # type: ignore

        if isinstance(value, np.generic):
            return _jsonable(value.item())
        if isinstance(value, np.ndarray):
            return [_jsonable(item) for item in value.tolist()]
    except Exception:
        pass
    try:
        import pandas as pd  # type: ignore

        if pd.isna(value):
            return None
    except Exception:
        pass
    return str(value)


def _parse_json(value: Any, default: Any) -> Any:
    if value in (None, ""):
        return default
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(str(value))
    except json.JSONDecodeError as exc:
        raise ToolError(f"invalid JSON argument: {exc}") from None


def _truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on"}


def _int(value: Any, default: int, *, minimum: int = 0, maximum: int = 10**9) -> int:
    try:
        number = int(float(str(value)))
    except (TypeError, ValueError):
        return default
    return max(minimum, min(maximum, number))


def _require_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"file not found: {path}")
    if not path.is_file():
        raise ToolError(f"not a regular file: {path}")


def _require_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _ensure_output(path: Path, overwrite: bool) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(
            f"{path} already exists; pass overwrite=true to replace it"
        )
    _require_parent(path)


def _verify_pdf(path: Path) -> int:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        _run(["pdfinfo", str(path)], check=True)
        text = subprocess.run(
            ["pdfinfo", str(path)], capture_output=True, text=True, timeout=30
        ).stdout
        match = re.search(r"Pages:\s+(\d+)", text)
        return int(match.group(1)) if match else 0
    reader = PdfReader(str(path))
    pages = len(reader.pages)
    if pages < 1:
        raise ToolError(f"generated PDF has no pages: {path}")
    return pages


def _run(
    argv: list[str],
    *,
    cwd: Path | None = None,
    timeout: int = 120,
    check: bool = False,
    input_text: str | None = None,
) -> subprocess.CompletedProcess:
    try:
        result = subprocess.run(
            argv,
            cwd=str(cwd) if cwd else None,
            capture_output=True,
            text=True,
            errors="replace",
            timeout=timeout,
            input=input_text,
        )
    except FileNotFoundError:
        raise ToolError(f"required program is not installed: {argv[0]}") from None
    except subprocess.TimeoutExpired:
        raise ToolError(f"command timed out after {timeout}s: {argv[0]}") from None
    if check and result.returncode != 0:
        detail = (result.stderr or result.stdout).strip()
        raise ToolError(f"{argv[0]} failed ({result.returncode}): {clip_detail(detail)}")
    return result


def clip_detail(text: str, limit: int = 4000) -> str:
    text = text.strip()
    return text if len(text) <= limit else text[-limit:]


# ---------------------------------------------------------------------------
# URL safety and HTTP
# ---------------------------------------------------------------------------
def _validate_url(url: str) -> urllib.parse.SplitResult:
    parsed = urllib.parse.urlsplit(str(url).strip())
    if parsed.scheme not in {"http", "https"}:
        raise ToolError("only http and https URLs are allowed")
    if not parsed.hostname:
        raise ToolError("URL has no host")
    if parsed.username or parsed.password:
        raise ToolError("URLs with embedded credentials are not allowed")
    host = parsed.hostname.strip("[]").lower()
    if host in {"localhost", "localhost.localdomain"} or host.endswith(".localhost"):
        raise ToolError("requests to localhost are not allowed")
    try:
        address = ipaddress.ip_address(host)
    except ValueError:
        address = None
    if address is not None and (
        address.is_private
        or address.is_loopback
        or address.is_link_local
        or address.is_multicast
        or address.is_reserved
        or address.is_unspecified
    ):
        raise ToolError("requests to private or loopback addresses are not allowed")
    return parsed


class _SafeRedirects(urllib.request.HTTPRedirectHandler):
    max_redirections = 5

    def redirect_request(self, req, fp, code, msg, headers, newurl):  # noqa: ANN001
        _validate_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


_OPENER = urllib.request.build_opener(_SafeRedirects())


def _http(
    method: str,
    url: str,
    *,
    headers: dict[str, str] | None = None,
    body: bytes | None = None,
    timeout: int = 30,
    max_bytes: int = MAX_HTTP_BYTES,
) -> dict[str, Any]:
    _validate_url(url)
    request = urllib.request.Request(
        url, data=body, method=method.upper(), headers=headers or {}
    )
    try:
        with _OPENER.open(request, timeout=timeout) as response:
            raw = response.read(max_bytes + 1)
            truncated = len(raw) > max_bytes
            raw = raw[:max_bytes]
            charset = response.headers.get_content_charset() or "utf-8"
            text = raw.decode(charset, "replace")
            safe_headers = {
                key: value
                for key, value in response.headers.items()
                if key.lower()
                not in {
                    "set-cookie",
                    "authorization",
                    "proxy-authorization",
                    "www-authenticate",
                }
            }
            return {
                "status": response.status,
                "reason": response.reason,
                "final_url": response.geturl(),
                "content_type": response.headers.get_content_type(),
                "charset": charset,
                "headers": safe_headers,
                "text": text,
                "truncated": truncated,
                "bytes": len(raw),
            }
    except urllib.error.HTTPError as exc:
        raw = exc.read(max_bytes) if exc.fp else b""
        text = raw.decode("utf-8", "replace")
        return {
            "status": exc.code,
            "reason": exc.reason,
            "final_url": url,
            "content_type": exc.headers.get_content_type() if exc.headers else "",
            "headers": {},
            "text": text,
            "truncated": False,
            "bytes": len(raw),
            "error": f"HTTP {exc.code}: {exc.reason}",
        }
    except urllib.error.URLError as exc:
        reason = getattr(exc, "reason", exc)
        raise ToolError(
            "network request failed (sandbox network policy may deny this "
            f"destination): {reason}"
        ) from None
    except TimeoutError:
        raise ToolError(f"request timed out after {timeout}s") from None


class _ReadableText(HTMLParser):
    """Extract readable text and the document title from HTML."""

    SKIP = {"script", "style", "noscript", "svg", "template"}
    BLOCK = {
        "p", "div", "br", "li", "tr", "section", "article", "header",
        "footer", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.title_parts: list[str] = []
        self._skip_depth = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in self.SKIP:
            self._skip_depth += 1
        if tag == "title":
            self._in_title = True
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self.SKIP and self._skip_depth:
            self._skip_depth -= 1
        if tag == "title":
            self._in_title = False
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title_parts.append(data)
            return
        if self._skip_depth:
            return
        self.parts.append(data)

    def text(self) -> str:
        raw = "".join(self.parts)
        lines = [re.sub(r"\s+", " ", line).strip() for line in raw.splitlines()]
        return "\n".join(line for line in lines if line)

    def title(self) -> str:
        return re.sub(r"\s+", " ", "".join(self.title_parts)).strip()


class _SearchResults(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.results: list[dict[str, str]] = []
        self._current: dict[str, str] | None = None
        self._capture: str | None = None

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        values = {key: value or "" for key, value in attrs}
        classes = values.get("class", "").split()
        if tag == "a" and "result__a" in classes:
            self._current = {"title": "", "url": values.get("href", ""), "snippet": ""}
            self._capture = "title"
        elif self._current is not None and "result__snippet" in classes:
            self._capture = "snippet"

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self._capture == "title":
            self._capture = None
        elif self._capture == "snippet" and tag in {"a", "div", "span"}:
            if self._current is not None:
                self.results.append(self._current)
            self._current = None
            self._capture = None

    def handle_data(self, data: str) -> None:
        if self._current is None or self._capture is None:
            return
        self._current[self._capture] += data

    def close(self) -> None:
        super().close()
        if self._current is not None:
            self.results.append(self._current)
            self._current = None


def _unwrap_duckduckgo(url: str) -> str:
    if "duckduckgo.com/l/" not in url:
        return url
    parsed = urllib.parse.urlsplit(url)
    query = urllib.parse.parse_qs(parsed.query)
    target = query.get("uddg", [""])[0]
    return urllib.parse.unquote(target) if target else url


# ---------------------------------------------------------------------------
# Web operations
# ---------------------------------------------------------------------------
def op_web_search(request: dict[str, Any]) -> dict[str, Any]:
    query = str(request.get("query", "")).strip()
    if not query:
        raise ToolError("query must not be empty")
    if len(query) > 500:
        raise ToolError("query is too long (maximum 500 characters)")
    limit = _int(request.get("max_results"), 5, minimum=1, maximum=10)
    data = urllib.parse.urlencode({"q": query, "kl": "wt-wt"}).encode()
    response = _http(
        "POST",
        "https://html.duckduckgo.com/html/",
        headers={
            "User-Agent": USER_AGENT,
            "Content-Type": "application/x-www-form-urlencoded",
        },
        body=data,
        timeout=_int(request.get("timeout"), 20, minimum=5, maximum=60),
    )
    parser = _SearchResults()
    parser.feed(response["text"])
    parser.close()
    results = []
    for item in parser.results:
        url = _unwrap_duckduckgo(item.get("url", ""))
        title = html.unescape(item.get("title", "")).strip()
        snippet = html.unescape(item.get("snippet", "")).strip()
        if not url or not title:
            continue
        results.append({"title": title, "url": url, "snippet": snippet})
        if len(results) >= limit:
            break
    return {
        "query": query,
        "backend": "duckduckgo-html",
        "status": response["status"],
        "results": results,
    }


def op_web_fetch(request: dict[str, Any]) -> dict[str, Any]:
    url = str(request.get("url", "")).strip()
    if not url:
        raise ToolError("url must not be empty")
    max_chars = _int(request.get("max_chars"), 12_000, minimum=200, maximum=MAX_TEXT_CHARS)
    response = _http(
        "GET",
        url,
        headers={"User-Agent": USER_AGENT},
        timeout=_int(request.get("timeout"), 20, minimum=5, maximum=60),
    )
    content_type = response.get("content_type", "")
    body = response.get("text", "")
    title = ""
    if "html" in content_type or body.lstrip().startswith("<"):
        parser = _ReadableText()
        parser.feed(body)
        parser.close()
        title = parser.title()
        text = parser.text()
    else:
        text = body
    if not text.strip():
        raise ToolError(
            f"no readable text extracted from {response.get('final_url', url)} "
            f"(status {response.get('status')})"
        )
    return {
        "url": url,
        "final_url": response.get("final_url"),
        "status": response.get("status"),
        "content_type": content_type,
        "title": title,
        "text": _clip(text, max_chars),
        "truncated": response.get("truncated", False) or len(text) > max_chars,
    }


def op_http_request(request: dict[str, Any]) -> dict[str, Any]:
    url = str(request.get("url", "")).strip()
    if not url:
        raise ToolError("url must not be empty")
    method = str(request.get("method", "GET")).upper()
    if method not in {"GET", "HEAD", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"}:
        raise ToolError(f"unsupported HTTP method: {method}")
    headers = _parse_json(request.get("headers"), {})
    if not isinstance(headers, dict):
        raise ToolError("headers must be a JSON object")
    headers = {str(key): str(value) for key, value in headers.items()}
    headers.setdefault("User-Agent", USER_AGENT)
    body_text = str(request.get("body", "") or "")
    body = body_text.encode("utf-8") if body_text and method not in {"GET", "HEAD"} else None
    response = _http(
        method,
        url,
        headers=headers,
        body=body,
        timeout=_int(request.get("timeout"), 30, minimum=5, maximum=120),
        max_bytes=_int(request.get("max_bytes"), MAX_HTTP_BYTES, minimum=1024, maximum=10_000_000),
    )
    response["request_method"] = method
    response["text"] = _clip(response.get("text", ""), _int(request.get("max_chars"), 12_000, minimum=200, maximum=MAX_TEXT_CHARS))
    return response


def op_download_url(request: dict[str, Any]) -> dict[str, Any]:
    url = str(request.get("url", "")).strip()
    destination = str(request.get("destination", "")).strip()
    if not url or not destination:
        raise ToolError("url and destination are required")
    overwrite = _truthy(request.get("overwrite"))
    max_bytes = _int(
        request.get("max_bytes"),
        MAX_DOWNLOAD_BYTES,
        minimum=1024,
        maximum=MAX_DOWNLOAD_BYTES,
    )
    workdir = str(request["workdir"])
    target = _resolve(workdir, destination)
    _ensure_output(target, overwrite)
    _validate_url(url)
    request_object = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    temp_path: Path | None = None
    try:
        with _OPENER.open(
            request_object,
            timeout=_int(request.get("timeout"), 60, minimum=5, maximum=300),
        ) as response:
            content_length = response.headers.get("Content-Length")
            if content_length and int(content_length) > max_bytes:
                raise ToolError(
                    f"download is larger than the {max_bytes}-byte limit"
                )
            size = 0
            with tempfile.NamedTemporaryFile(
                dir=target.parent, delete=False
            ) as handle:
                temp_path = Path(handle.name)
                while True:
                    chunk = response.read(64 * 1024)
                    if not chunk:
                        break
                    size += len(chunk)
                    if size > max_bytes:
                        raise ToolError(
                            f"download exceeded the {max_bytes}-byte limit"
                        )
                    handle.write(chunk)
            os.replace(temp_path, target)
            temp_path = None
            return {
                "url": url,
                "final_url": response.geturl(),
                "path": str(target),
                "bytes": size,
                "content_type": response.headers.get_content_type(),
                "status": response.status,
            }
    except urllib.error.HTTPError as exc:
        raise ToolError(f"download failed with HTTP {exc.code}: {exc.reason}") from None
    except urllib.error.URLError as exc:
        raise ToolError(
            "download failed (sandbox network policy may deny this destination): "
            f"{getattr(exc, 'reason', exc)}"
        ) from None
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink()
            except OSError:
                pass


# ---------------------------------------------------------------------------
# PDF and document operations
# ---------------------------------------------------------------------------
def _parse_pages(spec: str, total: int) -> list[int]:
    if not spec.strip():
        return list(range(1, total + 1))
    pages: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start, end = part.split("-", 1)
            first = int(start) if start else 1
            last = int(end) if end else total
            pages.update(range(max(1, first), min(total, last) + 1))
        else:
            number = int(part)
            if 1 <= number <= total:
                pages.add(number)
    if not pages:
        raise ToolError(f"no valid pages in {spec!r} for a {total}-page PDF")
    return sorted(pages)


def op_read_pdf(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    _require_file(path)
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        result = _run(
            ["pdftotext", str(path), "-"],
            timeout=_int(request.get("timeout"), 60, minimum=5, maximum=300),
        )
        if result.returncode != 0:
            raise ToolError(f"pdftotext failed: {clip_detail(result.stderr)}")
        return {
            "path": str(path),
            "pages": None,
            "metadata": {},
            "text": _clip(result.stdout, _int(request.get("max_chars"), 20_000, minimum=200, maximum=MAX_TEXT_CHARS)),
            "engine": "pdftotext",
        }
    reader = PdfReader(str(path))
    total = len(reader.pages)
    selected = _parse_pages(str(request.get("pages", "")), total)
    chunks = []
    for number in selected:
        text = reader.pages[number - 1].extract_text() or ""
        chunks.append(f"--- page {number} ---\n{text}")
    metadata = {}
    try:
        for key, value in (reader.metadata or {}).items():
            metadata[str(key)] = str(value)
    except Exception:
        metadata = {}
    return {
        "path": str(path),
        "pages": total,
        "selected_pages": selected,
        "metadata": metadata,
        "text": _clip("\n".join(chunks), _int(request.get("max_chars"), 20_000, minimum=200, maximum=MAX_TEXT_CHARS)),
        "engine": "pypdf",
    }


def _structured_content(raw: Any) -> dict[str, Any]:
    if isinstance(raw, dict):
        return raw
    value = str(raw or "")
    if value.lstrip().startswith("{"):
        try:
            parsed = json.loads(value)
            if isinstance(parsed, dict):
                return parsed
        except json.JSONDecodeError:
            pass
    return {"paragraphs": value.split("\n\n") if value else []}


def op_create_pdf(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    if path.suffix.lower() != ".pdf":
        raise ToolError("output path must end with .pdf")
    overwrite = _truthy(request.get("overwrite"))
    _ensure_output(path, overwrite)
    try:
        from reportlab.lib.pagesizes import A4  # type: ignore
        from reportlab.lib.styles import getSampleStyleSheet  # type: ignore
        from reportlab.lib.units import cm  # type: ignore
        from reportlab.platypus import (  # type: ignore
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError:
        raise ToolError(
            "reportlab is not installed in this sandbox image; use the "
            "approved application-tools image"
        ) from None
    content = _structured_content(request.get("content"))
    title = str(request.get("title", "") or content.get("title", ""))
    styles = getSampleStyleSheet()
    story = []
    if title:
        story.append(Paragraph(html.escape(title), styles["Title"]))
        story.append(Spacer(1, 0.4 * cm))
    for section in content.get("sections", []) or []:
        if not isinstance(section, dict):
            continue
        heading = str(section.get("heading", "")).strip()
        if heading:
            story.append(Paragraph(html.escape(heading), styles["Heading2"]))
        for paragraph in section.get("paragraphs", []) or []:
            story.append(Paragraph(html.escape(str(paragraph)), styles["BodyText"]))
            story.append(Spacer(1, 0.2 * cm))
        for bullet in section.get("bullets", []) or []:
            story.append(Paragraph("&bull; " + html.escape(str(bullet)), styles["BodyText"]))
    for paragraph in content.get("paragraphs", []) or []:
        story.append(Paragraph(html.escape(str(paragraph)), styles["BodyText"]))
        story.append(Spacer(1, 0.2 * cm))
    for table in content.get("tables", []) or []:
        rows = [[str(cell) for cell in row] for row in table]
        if rows:
            story.append(Table(rows, style=TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, "#888888"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ])))
            story.append(Spacer(1, 0.4 * cm))
    if not story:
        raise ToolError("no content supplied for the PDF")
    document = SimpleDocTemplate(
        str(path), pagesize=A4, title=title or path.stem,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
    )
    document.build(story)
    pages = _verify_pdf(path)
    return {"path": str(path), "pages": pages, "bytes": path.stat().st_size}


def op_edit_pdf(request: dict[str, Any]) -> dict[str, Any]:
    try:
        from pypdf import PdfReader, PdfWriter  # type: ignore
    except ImportError:
        raise ToolError("pypdf is not installed in this sandbox") from None
    workdir = str(request["workdir"])
    operation = str(request.get("operation", "")).strip().lower()
    overwrite = _truthy(request.get("overwrite"))
    if operation == "merge":
        inputs = _parse_json(request.get("inputs"), [])
        if not isinstance(inputs, list) or len(inputs) < 2:
            raise ToolError("merge requires a JSON list with at least two inputs")
        output = _resolve(workdir, str(request.get("output", "")))
        _ensure_output(output, overwrite)
        writer = PdfWriter()
        for item in inputs:
            source = _resolve(workdir, str(item))
            _require_file(source)
            reader = PdfReader(str(source))
            for page in reader.pages:
                writer.add_page(page)
        with output.open("wb") as handle:
            writer.write(handle)
        return {"operation": operation, "path": str(output), "pages": _verify_pdf(output)}

    source_value = request.get("input") or (request.get("inputs") or [""])[0]
    source = _resolve(workdir, str(source_value))
    _require_file(source)
    reader = PdfReader(str(source))
    total = len(reader.pages)

    if operation == "split":
        output_dir = _resolve(workdir, str(request.get("output", "split")))
        output_dir.mkdir(parents=True, exist_ok=True)
        written = []
        for number in range(1, total + 1):
            writer = PdfWriter()
            writer.add_page(reader.pages[number - 1])
            target = output_dir / f"{source.stem}-page-{number:03d}.pdf"
            if target.exists() and not overwrite:
                raise FileExistsError(
                    f"{target} already exists; pass overwrite=true"
                )
            with target.open("wb") as handle:
                writer.write(handle)
            written.append(str(target))
        return {"operation": operation, "directory": str(output_dir), "files": written}

    output = _resolve(workdir, str(request.get("output", "")))
    _ensure_output(output, overwrite)
    writer = PdfWriter()
    selected = _parse_pages(str(request.get("pages", "")), total)
    if operation == "extract":
        for number in selected:
            writer.add_page(reader.pages[number - 1])
    elif operation == "rotate":
        angle = _int(request.get("angle"), 90, minimum=-360, maximum=360)
        if angle % 90 != 0:
            raise ToolError("rotation angle must be a multiple of 90")
        writer.append(reader)
        for number in selected:
            writer.pages[number - 1].rotate(angle)
    else:
        raise ToolError(f"unsupported PDF operation: {operation!r}")
    with output.open("wb") as handle:
        writer.write(handle)
    return {"operation": operation, "path": str(output), "pages": _verify_pdf(output)}


def op_compile_latex(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    source = _resolve(workdir, str(request.get("path", "")))
    _require_file(source)
    if source.suffix.lower() != ".tex":
        raise ToolError("input must be a .tex file")
    output_value = str(request.get("output", "")).strip()
    output = (
        _resolve(workdir, output_value)
        if output_value
        else source.with_suffix(".pdf")
    )
    if output.suffix.lower() != ".pdf":
        raise ToolError("output must be a .pdf file")
    _ensure_output(output, _truthy(request.get("overwrite")))
    engine = str(request.get("engine", "pdflatex")).strip()
    if engine not in {"pdflatex", "xelatex", "lualatex", "latexmk"}:
        raise ToolError(f"unsupported LaTeX engine: {engine}")
    runs = _int(request.get("runs"), 2, minimum=1, maximum=5)
    timeout = _int(request.get("timeout"), 120, minimum=10, maximum=600)
    output_dir = output.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    command = [engine, "-interaction=nonstopmode", "-halt-on-error"]
    if engine == "latexmk":
        command.append("-pdf")
    command.extend([f"-output-directory={output_dir}", str(source)])
    diagnostics = []
    exit_code = 0
    for _ in range(runs):
        result = _run(command, cwd=source.parent, timeout=timeout)
        exit_code = result.returncode
        diagnostics.append((result.stdout or "") + (result.stderr or ""))
        if exit_code != 0:
            break
    produced = output_dir / f"{source.stem}.pdf"
    if exit_code != 0 or not produced.exists():
        log_path = output_dir / f"{source.stem}.log"
        log_text = ""
        if log_path.exists():
            log_text = log_path.read_text(encoding="utf-8", errors="replace")
        detail = clip_detail(log_text or diagnostics[-1] if diagnostics else "", 6000)
        raise ToolError(
            f"LaTeX compilation failed (exit {exit_code}). Diagnostics:\n{detail}"
        )
    if produced != output:
        if output.exists() and not _truthy(request.get("overwrite")):
            raise FileExistsError(f"{output} already exists; pass overwrite=true")
        produced.replace(output)
    pages = _verify_pdf(output)
    return {
        "path": str(output),
        "pages": pages,
        "engine": engine,
        "runs": len(diagnostics),
        "diagnostics": clip_detail(diagnostics[-1] if diagnostics else "", 2000),
    }


def op_read_docx(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    _require_file(path)
    if path.suffix.lower() != ".docx":
        raise ToolError("input must be a .docx file")
    try:
        from docx import Document  # type: ignore
    except ImportError:
        raise ToolError("python-docx is not installed in this sandbox") from None
    document = Document(str(path))
    paragraphs = [
        {"text": item.text, "style": item.style.name if item.style else ""}
        for item in document.paragraphs
        if item.text.strip()
    ]
    tables = []
    for table in document.tables:
        tables.append(
            [[cell.text for cell in row.cells] for row in table.rows]
        )
    return {
        "path": str(path),
        "paragraphs": paragraphs[:_int(request.get("max_paragraphs"), 500, minimum=1, maximum=5000)],
        "tables": tables[:_int(request.get("max_tables"), 20, minimum=1, maximum=200)],
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
    }


def op_write_docx(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    if path.suffix.lower() != ".docx":
        raise ToolError("output path must end with .docx")
    append = _truthy(request.get("append"))
    overwrite = _truthy(request.get("overwrite"))
    if path.exists() and not append and not overwrite:
        raise FileExistsError(f"{path} already exists; pass overwrite=true or append=true")
    _require_parent(path)
    try:
        from docx import Document  # type: ignore
    except ImportError:
        raise ToolError("python-docx is not installed in this sandbox") from None
    document = Document(str(path)) if append and path.exists() else Document()
    content = _structured_content(request.get("content"))
    title = str(request.get("title", "") or content.get("title", ""))
    if title:
        document.add_heading(title, level=0)
    for section in content.get("sections", []) or []:
        if not isinstance(section, dict):
            continue
        if section.get("heading"):
            document.add_heading(str(section["heading"]), level=1)
        for paragraph in section.get("paragraphs", []) or []:
            document.add_paragraph(str(paragraph))
        for bullet in section.get("bullets", []) or []:
            document.add_paragraph(str(bullet), style="List Bullet")
    for paragraph in content.get("paragraphs", []) or []:
        document.add_paragraph(str(paragraph))
    for table in content.get("tables", []) or []:
        rows = list(table)
        if not rows:
            continue
        width = max(len(row) for row in rows)
        table_object = document.add_table(rows=len(rows), cols=width)
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table_object.cell(row_index, column_index).text = str(value)
    document.save(str(path))
    if not path.exists() or path.stat().st_size == 0:
        raise ToolError("DOCX was not created")
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
    }


def op_read_xlsx(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    _require_file(path)
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ToolError("input must be an .xlsx or .xlsm file")
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        raise ToolError("openpyxl is not installed in this sandbox") from None
    workbook = load_workbook(str(path), read_only=True, data_only=True)
    requested = str(request.get("sheet", "")).strip()
    sheets = [requested] if requested else list(workbook.sheetnames)
    max_rows = _int(request.get("max_rows"), 200, minimum=1, maximum=5000)
    max_cols = _int(request.get("max_cols"), 50, minimum=1, maximum=500)
    result = {}
    for name in sheets:
        if name not in workbook.sheetnames:
            raise ToolError(f"unknown worksheet: {name!r}")
        worksheet = workbook[name]
        columns = min(max_cols, max(1, worksheet.max_column or 1))
        rows = []
        for row in worksheet.iter_rows(
            max_row=max_rows, max_col=columns, values_only=True
        ):
            values = [_jsonable(value) for value in row]
            while values and values[-1] is None:
                values.pop()
            rows.append(values)
        result[name] = {
            "rows": rows,
            "max_row": worksheet.max_row,
            "max_column": worksheet.max_column,
            "truncated": (worksheet.max_row or 0) > max_rows,
        }
    workbook.close()
    return {"path": str(path), "sheets": result}


def op_write_xlsx(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ToolError("output path must end with .xlsx")
    append = _truthy(request.get("append"))
    overwrite = _truthy(request.get("overwrite"))
    if path.exists() and not append and not overwrite:
        raise FileExistsError(f"{path} already exists; pass overwrite=true or append=true")
    _require_parent(path)
    try:
        from openpyxl import Workbook, load_workbook  # type: ignore
    except ImportError:
        raise ToolError("openpyxl is not installed in this sandbox") from None
    workbook = load_workbook(str(path)) if append and path.exists() else Workbook()
    content = _structured_content(request.get("content"))
    datasheets = content.get("sheets") or []
    if not datasheets:
        datasheets = [{
            "name": str(request.get("sheet", "Sheet1")),
            "rows": content.get("rows", []),
        }]
    for index, sheet in enumerate(datasheets):
        if not isinstance(sheet, dict):
            continue
        name = str(sheet.get("name") or f"Sheet{index + 1}")
        if name in workbook.sheetnames:
            worksheet = workbook[name]
        else:
            worksheet = workbook.create_sheet(title=name)
        rows = sheet.get("rows", [])
        for row in rows:
            worksheet.append([_jsonable(value) for value in row])
    default_sheet = workbook["Sheet"] if "Sheet" in workbook.sheetnames else None
    if default_sheet is not None and len(workbook.sheetnames) > 1 and default_sheet.max_row == 1:
        workbook.remove(default_sheet)
    workbook.save(str(path))
    if not path.exists() or path.stat().st_size == 0:
        raise ToolError("XLSX was not created")
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sheets": workbook.sheetnames,
    }


def op_create_pptx(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    if path.suffix.lower() != ".pptx":
        raise ToolError("output path must end with .pptx")
    _ensure_output(path, _truthy(request.get("overwrite")))
    try:
        from pptx import Presentation  # type: ignore
    except ImportError:
        raise ToolError("python-pptx is not installed in this sandbox") from None
    content = _structured_content(request.get("content"))
    presentation = Presentation()
    title = str(request.get("title", "") or content.get("title", ""))
    slides = content.get("slides") or []
    if not slides:
        paragraphs = content.get("paragraphs", [])
        if not paragraphs:
            raise ToolError("no slide content supplied")
        slides = [{"title": title, "bullets": paragraphs}]
    for index, slide in enumerate(slides):
        if not isinstance(slide, dict):
            continue
        layout = presentation.slide_layouts[0] if index == 0 and title else presentation.slide_layouts[1]
        page = presentation.slides.add_slide(layout)
        if page.shapes.title is not None:
            page.shapes.title.text = str(slide.get("title") or (title if index == 0 else f"Slide {index + 1}"))
        body = None
        for placeholder in page.placeholders:
            if placeholder.placeholder_format.idx == 1:
                body = placeholder
                break
        if body is not None:
            frame = body.text_frame
            frame.clear()
            bullets = [str(item) for item in slide.get("bullets", [])]
            if bullets:
                frame.text = bullets[0]
                for bullet in bullets[1:]:
                    frame.add_paragraph().text = bullet
    presentation.save(str(path))
    if not path.exists() or path.stat().st_size == 0:
        raise ToolError("PPTX was not created")
    return {"path": str(path), "bytes": path.stat().st_size, "slides": len(presentation.slides)}


# ---------------------------------------------------------------------------
# Database and dataset operations
# ---------------------------------------------------------------------------
_SQL_READ = {"SELECT", "WITH", "PRAGMA", "EXPLAIN", "VALUES"}
_SQL_WRITE = re.compile(
    r"\b(INSERT|UPDATE|DELETE|CREATE|DROP|ALTER|REPLACE|VACUUM|ATTACH|DETACH|TRUNCATE|GRANT|REVOKE)\b",
    re.IGNORECASE,
)


def is_remote_database(database: str) -> bool:
    return "://" in database


def sql_is_read_only(query: str) -> bool:
    cleaned = re.sub(r"--[^\n]*", " ", query)
    cleaned = re.sub(r"/\*.*?\*/", " ", cleaned, flags=re.DOTALL).strip()
    if not cleaned:
        return False
    first = cleaned.split(None, 1)[0].upper()
    if first not in _SQL_READ:
        return False
    if first == "WITH" and _SQL_WRITE.search(cleaned):
        return False
    return True


def _sqlite_path(workdir: str, database: str, *, create: bool = False) -> Path:
    if database in {":memory:", ""}:
        raise ToolError("a workspace database file path is required")
    path = _resolve(workdir, database)
    if not path.exists():
        if not create:
            raise FileNotFoundError(f"database not found: {path}")
        _require_parent(path)
        return path
    if not path.is_file():
        raise ToolError(f"not a database file: {path}")
    return path


def _sqlite_rows(cursor: sqlite3.Cursor, max_rows: int) -> dict[str, Any]:
    if cursor.description is None:
        return {"columns": [], "rows": [], "rowcount": cursor.rowcount}
    columns = [item[0] for item in cursor.description]
    rows = cursor.fetchmany(max_rows + 1)
    truncated = len(rows) > max_rows
    return {
        "columns": columns,
        "rows": [[_jsonable(value) for value in row] for row in rows[:max_rows]],
        "rowcount": len(rows[:max_rows]),
        "truncated": truncated,
    }


def op_sql_query(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    database = str(request.get("database", "")).strip()
    query = str(request.get("query", "")).strip()
    if not database or not query:
        raise ToolError("database and query are required")
    if query.rstrip().endswith(";"):
        query = query.rstrip()[:-1]
    read_only = _truthy(request.get("read_only"))
    max_rows = _int(request.get("max_rows"), 200, minimum=1, maximum=MAX_ROWS)
    params = _parse_json(request.get("params"), [])
    if params in ({}, None):
        params = []
    if not isinstance(params, (list, tuple, dict)):
        raise ToolError("params must be a JSON array or object")
    if read_only and not sql_is_read_only(query):
        raise ToolError("read_only=true was requested but the query modifies data")
    if is_remote_database(database):
        try:
            from sqlalchemy import create_engine, text  # type: ignore
        except ImportError:
            raise ToolError(
                "non-SQLite databases require SQLAlchemy and an installed driver "
                "in the sandbox image"
            ) from None
        engine = create_engine(database, future=True)
        with engine.connect() as connection:
            result = connection.execute(text(query), params if isinstance(params, dict) else tuple(params))
            if result.returns_rows:
                columns = list(result.keys())
                rows = result.fetchmany(max_rows + 1)
                return {
                    "engine": engine.dialect.name,
                    "columns": columns,
                    "rows": [[_jsonable(value) for value in row] for row in rows[:max_rows]],
                    "rowcount": min(len(rows), max_rows),
                    "truncated": len(rows) > max_rows,
                    "read_only": sql_is_read_only(query),
                }
            if not read_only:
                connection.commit()
            return {
                "engine": engine.dialect.name,
                "columns": [],
                "rows": [],
                "rowcount": result.rowcount,
                "read_only": False,
            }
    path = _sqlite_path(workdir, database, create=not read_only)
    connection = sqlite3.connect(str(path), timeout=30)
    try:
        if read_only:
            connection.execute("PRAGMA query_only = ON")
        cursor = connection.execute(query, params)
        result = _sqlite_rows(cursor, max_rows)
        if not read_only:
            connection.commit()
        result["engine"] = "sqlite"
        result["read_only"] = sql_is_read_only(query)
        return result
    except sqlite3.Error as exc:
        raise ToolError(f"SQL error: {exc}") from None
    finally:
        connection.close()


def op_database_schema(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    database = str(request.get("database", "")).strip()
    if not database:
        raise ToolError("database is required")
    if is_remote_database(database):
        try:
            from sqlalchemy import create_engine, inspect  # type: ignore
        except ImportError:
            raise ToolError("SQLAlchemy is required for non-SQLite databases") from None
        engine = create_engine(database, future=True)
        inspector = inspect(engine)
        tables = inspector.get_table_names()
        requested = str(request.get("table", "")).strip()
        if requested and requested not in tables:
            raise ToolError(f"unknown table: {requested!r}")
        result = {}
        for table in ([requested] if requested else tables):
            result[table] = {
                "columns": inspector.get_columns(table),
                "indexes": inspector.get_indexes(table),
                "foreign_keys": inspector.get_foreign_keys(table),
            }
        return {"engine": engine.dialect.name, "tables": result}
    path = _sqlite_path(workdir, database)
    connection = sqlite3.connect(str(path), timeout=30)
    try:
        requested = str(request.get("table", "")).strip()
        names = [
            row[0]
            for row in connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            )
        ]
        if requested and requested not in names:
            raise ToolError(f"unknown table: {requested!r}")
        result = {}
        for table in ([requested] if requested else names):
            quoted = table.replace('"', '""')
            columns = [
                {
                    "name": row[1],
                    "type": row[2],
                    "nullable": not bool(row[3]),
                    "default": row[4],
                    "primary_key": bool(row[5]),
                }
                for row in connection.execute(f'PRAGMA table_info("{quoted}")')
            ]
            indexes = []
            for row in connection.execute(f'PRAGMA index_list("{quoted}")'):
                index_name = row[1]
                index_quoted = str(index_name).replace('"', '""')
                index_columns = [
                    info[2]
                    for info in connection.execute(
                        f'PRAGMA index_info("{index_quoted}")'
                    )
                ]
                indexes.append(
                    {
                        "name": index_name,
                        "unique": bool(row[2]),
                        "columns": index_columns,
                    }
                )
            foreign_keys = [
                {
                    "id": row[0],
                    "from": row[3],
                    "table": row[2],
                    "to": row[4],
                }
                for row in connection.execute(f'PRAGMA foreign_key_list("{quoted}")')
            ]
            result[table] = {
                "columns": columns,
                "indexes": indexes,
                "foreign_keys": foreign_keys,
            }
        return {"engine": "sqlite", "database": str(path), "tables": result}
    except sqlite3.Error as exc:
        raise ToolError(f"SQL error: {exc}") from None
    finally:
        connection.close()


def _import_pandas():
    try:
        import pandas as pd  # type: ignore
    except ImportError:
        raise ToolError("pandas is not installed in this sandbox") from None
    return pd


def op_read_data(request: dict[str, Any]) -> dict[str, Any]:
    pd = _import_pandas()
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    _require_file(path)
    file_format = str(request.get("format", "")).strip().lower() or path.suffix.lower().lstrip(".")
    max_rows = _int(request.get("max_rows"), 200, minimum=1, maximum=MAX_ROWS)
    if file_format in {"csv", "tsv", "txt"}:
        separator = "\t" if file_format == "tsv" else ("," if file_format != "txt" else None)
        frame = pd.read_csv(str(path), sep=separator, engine="python")
    elif file_format == "json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            payload = payload.get("data", payload.get("records", payload))
        frame = pd.DataFrame(payload)
    elif file_format == "jsonl":
        frame = pd.read_json(str(path), lines=True)
    elif file_format == "parquet":
        frame = pd.read_parquet(str(path))
    elif file_format in {"xlsx", "xlsm"}:
        frame = pd.read_excel(str(path), sheet_name=str(request.get("sheet", 0) or 0))
    else:
        raise ToolError(f"unsupported data format: {file_format!r}")
    columns = [str(item) for item in frame.columns]
    rows = frame.head(max_rows).to_dict(orient="records")
    return {
        "path": str(path),
        "format": file_format,
        "columns": columns,
        "dtypes": {str(key): str(value) for key, value in frame.dtypes.items()},
        "rows": [_jsonable(row) for row in rows],
        "shape": list(frame.shape),
        "truncated": len(frame) > max_rows,
    }


def op_write_data(request: dict[str, Any]) -> dict[str, Any]:
    pd = _import_pandas()
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    _ensure_output(path, _truthy(request.get("overwrite")))
    file_format = str(request.get("format", "")).strip().lower() or path.suffix.lower().lstrip(".")
    data = _parse_json(request.get("data"), [])
    if isinstance(data, dict) and "columns" in data and "rows" in data:
        frame = pd.DataFrame(data["rows"], columns=data["columns"])
    else:
        frame = pd.DataFrame(data)
    index = _truthy(request.get("index"))
    if file_format in {"csv", "tsv"}:
        frame.to_csv(str(path), sep="\t" if file_format == "tsv" else ",", index=index)
    elif file_format == "json":
        frame.to_json(str(path), orient="records", indent=2, force_ascii=False)
    elif file_format == "jsonl":
        frame.to_json(str(path), orient="records", lines=True, force_ascii=False)
    elif file_format == "parquet":
        frame.to_parquet(str(path), index=index)
    elif file_format in {"xlsx", "xlsm"}:
        frame.to_excel(str(path), index=index)
    else:
        raise ToolError(f"unsupported data format: {file_format!r}")
    if not path.exists() or path.stat().st_size == 0:
        raise ToolError("output dataset was not created")
    return {
        "path": str(path),
        "format": file_format,
        "rows": int(len(frame)),
        "columns": [str(item) for item in frame.columns],
        "bytes": path.stat().st_size,
    }


NOTEBOOK_RUNNER = r'''
import ast
import contextlib
import io
import json
import sys
import traceback

notebook_path, result_path, allow_errors = sys.argv[1], sys.argv[2], sys.argv[3] == "1"
with open(notebook_path, "r", encoding="utf-8") as handle:
    notebook = json.load(handle)
namespace = {"__name__": "__main__"}
results = []
failed = False
for index, cell in enumerate(notebook.get("cells", []), start=1):
    if cell.get("cell_type") != "code":
        continue
    source = cell.get("source", "")
    if isinstance(source, list):
        source = "".join(source)
    stdout = io.StringIO()
    stderr = io.StringIO()
    error = None
    try:
        with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
            tree = ast.parse(source, filename=f"<cell {index}>")
            body = tree.body
            if body and isinstance(body[-1], ast.Expr):
                if body[:-1]:
                    exec(compile(ast.Module(body=body[:-1], type_ignores=[]), f"<cell {index}>", "exec"), namespace)
                value = eval(compile(ast.Expression(body[-1].value), f"<cell {index}>", "eval"), namespace)
                if value is not None:
                    print(repr(value))
            else:
                exec(compile(tree, f"<cell {index}>", "exec"), namespace)
    except Exception:
        error = traceback.format_exc()
    results.append({
        "cell": index,
        "stdout": stdout.getvalue(),
        "stderr": stderr.getvalue(),
        "error": error,
    })
    if error and not allow_errors:
        failed = True
        break
with open(result_path, "w", encoding="utf-8") as handle:
    json.dump({"cells": results, "failed": failed}, handle)
'''


def _zmq_works() -> bool:
    """Probe whether a Jupyter kernel can start in this sandbox.

    Some sandbox process policies deny the thread-scheduling call that ZeroMQ
    makes during teardown. A failed probe means the direct Python runner is
    used instead of risking a killed worker process.
    """
    probe_code = (
        "import nbformat\n"
        "from nbclient import NotebookClient\n"
        "nb = nbformat.v4.new_notebook("
        "cells=[nbformat.v4.new_code_cell('print(1)')])\n"
        "NotebookClient(nb, timeout=20).execute()\n"
        "print('ok')\n"
    )
    try:
        probe = subprocess.run(
            [sys.executable, "-c", probe_code],
            capture_output=True,
            text=True,
            timeout=40,
        )
    except (OSError, subprocess.TimeoutExpired):
        return False
    return probe.returncode == 0 and "ok" in probe.stdout


def _execute_notebook_fallback(
    notebook: Any,
    *,
    cwd: Path,
    timeout: int,
    allow_errors: bool,
) -> tuple[list[dict[str, Any]], bool]:
    """Execute code cells with the sandbox Python when Jupyter cannot start."""
    import nbformat  # type: ignore

    with tempfile.TemporaryDirectory(dir=cwd) as directory:
        root = Path(directory)
        source_path = root / "notebook.json"
        result_path = root / "result.json"
        runner_path = root / "runner.py"
        source_path.write_text(
            nbformat.writes(notebook), encoding="utf-8"
        )
        runner_path.write_text(NOTEBOOK_RUNNER, encoding="utf-8")
        code_cells = sum(
            1 for cell in notebook.cells if cell.cell_type == "code"
        )
        total_timeout = min(1800, timeout * max(1, code_cells) + 30)
        process = subprocess.run(
            [
                sys.executable,
                str(runner_path),
                str(source_path),
                str(result_path),
                "1" if allow_errors else "0",
            ],
            cwd=str(cwd),
            capture_output=True,
            text=True,
            errors="replace",
            timeout=total_timeout,
        )
        if not result_path.exists():
            raise ToolError(
                "notebook execution failed: "
                + clip_detail(process.stderr or process.stdout)
            )
        payload = json.loads(result_path.read_text(encoding="utf-8"))
    return list(payload.get("cells", [])), bool(payload.get("failed"))


def _apply_notebook_results(
    notebook: Any,
    results: list[dict[str, Any]],
    *,
    allow_errors: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Copy runner results back into notebook cells and build a summary."""
    import nbformat  # type: ignore

    summary_cells: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    by_index = {item["cell"]: item for item in results}
    execution = 0
    for index, cell in enumerate(notebook.cells, start=1):
        if cell.cell_type != "code":
            continue
        result = by_index.get(index)
        if result is None:
            continue
        execution += 1
        cell["execution_count"] = execution
        cell["outputs"] = []
        if result.get("stdout"):
            cell["outputs"].append(
                {
                    "output_type": "stream",
                    "name": "stdout",
                    "text": result["stdout"],
                }
            )
        if result.get("stderr"):
            cell["outputs"].append(
                {
                    "output_type": "stream",
                    "name": "stderr",
                    "text": result["stderr"],
                }
            )
        if result.get("error"):
            traceback_text = str(result["error"])
            last = [line for line in traceback_text.strip().splitlines() if line]
            message = last[-1] if last else "execution failed"
            name, _, value = message.partition(":")
            cell["outputs"].append(
                {
                    "output_type": "error",
                    "ename": name.strip() or "Error",
                    "evalue": value.strip(),
                    "traceback": traceback_text.splitlines(),
                }
            )
            errors.append({"cell": index, "error": message.strip()})
        cell["outputs"] = [
            nbformat.from_dict(output) for output in cell["outputs"]
        ]
        output = "\n".join(
            part
            for part in (result.get("stdout"), result.get("stderr"), result.get("error"))
            if part
        )
        summary_cells.append(
            {
                "cell": index,
                "execution_count": execution,
                "output": _clip(output, 2000),
            }
        )
        if result.get("error") and not allow_errors:
            break
    return summary_cells, errors


def op_execute_notebook(request: dict[str, Any]) -> dict[str, Any]:
    workdir = str(request["workdir"])
    path = _resolve(workdir, str(request.get("path", "")))
    _require_file(path)
    if path.suffix.lower() != ".ipynb":
        raise ToolError("input must be a .ipynb notebook")
    try:
        import nbformat  # type: ignore
    except ImportError:
        raise ToolError("nbformat is not installed in this sandbox") from None
    timeout = _int(request.get("timeout"), 120, minimum=10, maximum=1800)
    allow_errors = _truthy(request.get("allow_errors"))
    destination_value = str(request.get("output", "")).strip()
    destination = (
        _resolve(workdir, destination_value)
        if destination_value
        else None
    )
    if destination is not None:
        _ensure_output(destination, _truthy(request.get("overwrite")))
    notebook = nbformat.read(str(path), as_version=4)

    cells: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    engine = "subprocess"
    if _zmq_works():
        try:
            from nbclient import NotebookClient  # type: ignore

            client = NotebookClient(
                notebook,
                timeout=timeout,
                kernel_name=str(request.get("kernel", "python3")),
                allow_errors=allow_errors,
                resources={"metadata": {"path": str(path.parent)}},
            )
            client.execute()
            engine = "nbclient"
        except Exception as exc:
            # A kernel failure must not lose the notebook. Fall back to direct
            # sandbox execution and keep the original error as diagnostics.
            errors.append({"engine": "nbclient", "error": str(exc)})
        if engine == "nbclient":
            for index, cell in enumerate(notebook.cells, start=1):
                if cell.cell_type != "code":
                    continue
                summary = []
                for output in cell.get("outputs", []):
                    if output.get("output_type") == "stream":
                        summary.append(str(output.get("text", "")))
                    elif output.get("output_type") in {"execute_result", "display_data"}:
                        summary.append(str(output.get("data", {}).get("text/plain", "")))
                    elif output.get("output_type") == "error":
                        message = f"{output.get('ename')}: {output.get('evalue')}"
                        summary.append(message)
                        errors.append({"cell": index, "error": message})
                cells.append(
                    {
                        "cell": index,
                        "execution_count": cell.get("execution_count"),
                        "output": _clip("\n".join(part for part in summary if part), 2000),
                    }
                )

    if engine != "nbclient":
        results, failed = _execute_notebook_fallback(
            notebook,
            cwd=path.parent,
            timeout=timeout,
            allow_errors=allow_errors,
        )
        cells, fallback_errors = _apply_notebook_results(
            notebook, results, allow_errors=allow_errors
        )
        errors.extend(fallback_errors)
        if failed:
            errors.append({"engine": "subprocess", "error": "execution stopped on first error"})

    if destination is not None:
        nbformat.write(notebook, str(destination))
    return {
        "path": str(path),
        "output": str(destination) if destination is not None else None,
        "engine": engine,
        "executed_cells": len(cells),
        "errors": errors,
        "cells": cells[:50],
    }


OPERATIONS = {
    "web_search": op_web_search,
    "web_fetch": op_web_fetch,
    "http_request": op_http_request,
    "download_url": op_download_url,
    "read_pdf": op_read_pdf,
    "create_pdf": op_create_pdf,
    "edit_pdf": op_edit_pdf,
    "compile_latex": op_compile_latex,
    "read_docx": op_read_docx,
    "write_docx": op_write_docx,
    "read_xlsx": op_read_xlsx,
    "write_xlsx": op_write_xlsx,
    "create_pptx": op_create_pptx,
    "sql_query": op_sql_query,
    "database_schema": op_database_schema,
    "read_data": op_read_data,
    "write_data": op_write_data,
    "execute_notebook": op_execute_notebook,
}


def handle(request: dict[str, Any]) -> dict[str, Any]:
    operation = str(request.get("op", ""))
    handler = OPERATIONS.get(operation)
    if handler is None:
        raise ValueError(f"unsupported operation: {operation!r}")
    return handler(request)


def main() -> int:
    try:
        raw = sys.stdin.read()
        if not raw.strip():
            raise ValueError("empty tool request")
        request = json.loads(raw)
        if not isinstance(request, dict):
            raise ValueError("tool request must be a JSON object")
        response = {"ok": True, **handle(request)}
    except Exception as exc:
        response = {
            "ok": False,
            "error": {
                "type": type(exc).__name__,
                "message": str(exc),
            },
        }
    sys.stdout.write(json.dumps(response, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
